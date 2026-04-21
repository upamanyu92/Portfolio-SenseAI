from io import BytesIO
import unittest

import openpyxl
from docx import Document

from app.parser import extract_text


class ParserTests(unittest.TestCase):
    def test_extract_text_from_xlsx(self) -> None:
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(["Ticker", "Qty"])
        sheet.append(["TCS", 10])

        file_obj = BytesIO()
        workbook.save(file_obj)
        file_obj.seek(0)

        extracted = extract_text(file_obj.getvalue(), ".xlsx")
        self.assertIn("Ticker,Qty", extracted)
        self.assertIn("TCS,10", extracted)

    def test_extract_text_from_docx(self) -> None:
        document = Document()
        document.add_paragraph("INFY 25")
        file_obj = BytesIO()
        document.save(file_obj)
        file_obj.seek(0)

        extracted = extract_text(file_obj.getvalue(), ".docx")
        self.assertEqual(extracted, "INFY 25")

    def test_extract_text_for_unsupported_extension(self) -> None:
        self.assertEqual(extract_text(b"abc", ".csv"), "")


if __name__ == "__main__":
    unittest.main()
